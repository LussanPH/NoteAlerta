from io import BytesIO
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from models import Notificacao

MEDIA_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMATO_DATA = "DD/MM/YYYY HH:MM"

FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", start_color="0F6E56", end_color="0F6E56")
FUNDO_CAMPO = PatternFill("solid", start_color="D3D3D3", end_color="D3D3D3")


# Escreve texto sempre como string: o openpyxl trata "=..." como fórmula, e esses campos vêm de usuários
def _escrever(ws, linha: int, coluna: int, valor):
    if valor is None or valor == "":
        valor = "-"
    celula = ws.cell(row=linha, column=coluna)

    if isinstance(valor, str):
        celula.value = ILLEGAL_CHARACTERS_RE.sub("", valor)
        celula.data_type = "s"
    else:
        celula.value = valor

    return celula


def _larguras(ws, minimo: int = 10, maximo: int = 60):
    for coluna in ws.columns:
        maior = max(len(str(c.value)) for c in coluna if c.value is not None)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(maior + 2, minimo), maximo)


# Serializa a planilha e devolve como download (equivalente ao antigo retorno do PDF)
def xlsx_response(wb: Workbook, nome_arquivo: str) -> Response:
    buffer = BytesIO()
    wb.save(buffer)

    return Response(
        content=buffer.getvalue(),
        media_type=MEDIA_TYPE_XLSX,
        headers={
            "Content-Disposition": f'attachment; filename="{nome_arquivo}"'
        }
    )


# notificacoes: lista de Notificacao ou de tuplas (Notificacao, *extras).
# colunas_extras: título das colunas dos valores extras da tupla, na mesma ordem.
def relatorio_consolidado(cabecalho: dict, notificacoes: list, colunas_extras: tuple[str, ...] = ()) -> Workbook:
    wb = Workbook()

    ws = wb.active
    ws.title = "Notificações"

    titulos = ["ID", "Categoria", "Evento", "Status", "Local", "Data", *colunas_extras]
    for col, titulo in enumerate(titulos, start=1):
        c = _escrever(ws, 1, col, titulo)
        c.font = FONTE_CABECALHO
        c.fill = FUNDO_CABECALHO

    for linha, item in enumerate(notificacoes, start=2):
        n, *extras = (item,) if isinstance(item, Notificacao) else item

        _escrever(ws, linha, 1, n.id)
        _escrever(ws, linha, 2, n.categoria)
        _escrever(ws, linha, 3, n.tipo_evento)
        _escrever(ws, linha, 4, n.status)
        _escrever(ws, linha, 5, n.local_ocorrencia)
        data = _escrever(ws, linha, 6, n.data_ocorrencia)
        data.number_format = FORMATO_DATA
        for col, valor in enumerate(extras, start=7):
            _escrever(ws, linha, col, valor)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _larguras(ws)

    resumo = wb.create_sheet("Resumo")
    campos = {**cabecalho, "Total de notificações": len(notificacoes)}
    for linha, (campo, valor) in enumerate(campos.items(), start=1):
        _escrever(resumo, linha, 1, campo).font = Font(bold=True)
        _escrever(resumo, linha, 2, valor)
    _larguras(resumo)

    return wb


def relatorio_notificacao(notificacao: Notificacao) -> Workbook:
    wb = Workbook()

    ws = wb.active
    ws.title = "Notificação"

    campos = [
        ("Protocolo", f"#{str(notificacao.id).zfill(7)}"),
        ("Município", notificacao.municipio),
        ("Tipo do evento", notificacao.tipo_evento),
        ("Categoria", notificacao.categoria),
        ("Status", notificacao.status),
        ("Data de ocorrência", notificacao.data_ocorrencia),
        ("Local", notificacao.local_ocorrencia),
        ("Afetados", notificacao.pessoas_animais_infectados_afetados or 0),
        ("Continuidade da situação", notificacao.continuidade_situacao),
        ("Descrição", notificacao.descricao or "Nenhuma descrição informada."),
    ]

    for linha, (campo, valor) in enumerate(campos, start=1):
        c = _escrever(ws, linha, 1, campo)
        c.font = Font(bold=True)
        c.fill = FUNDO_CAMPO
        v = _escrever(ws, linha, 2, valor)
        v.alignment = Alignment(wrap_text=True, vertical="top")
        if campo == "Data de ocorrência":
            v.number_format = FORMATO_DATA

    _larguras(ws)

    return wb
